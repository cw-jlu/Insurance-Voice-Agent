"""Extract and clean the competition data under data/raw without altering any content.

Invariant enforced for every text output: removing all whitespace from the cleaned text
yields exactly the same string as removing all whitespace from the raw extraction.
"""
import hashlib, json, re, sys
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.utils import range_boundaries
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/尊享百万"
EXTRACTED = ROOT / "data/extracted"
CLEAN = ROOT / "data/clean"
WS = re.compile(r"\s+")
CJK = r"　-〿一-鿿＀-￯“”‘’"
report = {"pdf": {}, "xlsx": {}, "md": {}}


def strip_ws(s):
    return WS.sub("", s)


def clean_pdf_text(text):
    # Rejoin hard-wrapped lines, then drop layout spaces that sit next to CJK characters.
    s = text.replace("\r", "").replace("\n", "")
    s = s.replace(" ", " ")
    s = re.sub(rf"(?<=[{CJK}])[ \t]+|[ \t]+(?=[{CJK}])", "", s)
    s = re.sub(r"[ \t]{2,}", " ", s)
    # Put each enumerated clause / numbered item on its own line for readability (insertion of \n only).
    s = re.sub(r"(?<=[。；：\d])(?=[一二三四五六七八九十]{1,3}、)", "\n", s)
    s = re.sub(r"(?<=[；。：])(?=\d{1,2}[\.、)）](?!\d))", "\n", s)
    s = re.sub(r"(?<=[；。：])(?=[（(][一二三四五六七八九十\d]+[）)])", "\n", s)
    return s.strip() + "\n"


def extract_pdfs():
    (EXTRACTED / "pdf_text").mkdir(parents=True, exist_ok=True)
    (CLEAN / "text").mkdir(parents=True, exist_ok=True)
    for pdf in sorted(RAW.glob("*.pdf")):
        stem = pdf.stem
        with pdfplumber.open(pdf) as p:
            plumber_pages = [pg.extract_text() or "" for pg in p.pages]
        pypdf_pages = [pg.extract_text() or "" for pg in PdfReader(pdf).pages]
        raw = "".join(f"<<<PAGE {i + 1}>>>\n{t}\n" for i, t in enumerate(plumber_pages))
        (EXTRACTED / "pdf_text" / f"{stem}.pdfplumber.txt").write_text(raw, encoding="utf-8")
        (EXTRACTED / "pdf_text" / f"{stem}.pypdf.txt").write_text(
            "".join(f"<<<PAGE {i + 1}>>>\n{t}\n" for i, t in enumerate(pypdf_pages)), encoding="utf-8")

        cross = [strip_ws(a) == strip_ws(b) for a, b in zip(plumber_pages, pypdf_pages)]
        cleaned = "".join(f"<!-- page {i + 1} -->\n{clean_pdf_text(t)}\n" for i, t in enumerate(plumber_pages))
        body = re.sub(r"<!-- page \d+ -->", "", cleaned)
        assert strip_ws(body) == strip_ws("".join(plumber_pages)), f"cleaning altered content: {stem}"
        (CLEAN / "text" / f"{stem}.md").write_text(
            f"<!-- source: data/raw/尊享百万/{pdf.name}; sha256={hashlib.sha256(pdf.read_bytes()).hexdigest()} -->\n"
            f"<!-- cleaned: line-wrap/space removal only; non-whitespace chars identical to raw extraction -->\n\n"
            + cleaned, encoding="utf-8")
        report["pdf"][pdf.name] = {
            "pages": len(plumber_pages),
            "chars_non_ws": len(strip_ws("".join(plumber_pages))),
            "pypdf_vs_pdfplumber_equal_per_page": cross,
            "clean_equals_raw_non_ws": True,
        }


def cell_text(c):
    v = c.value
    if v is None:
        return None
    if isinstance(v, (int, float)) and "%" in (c.number_format or ""):
        return f"{v * 100:g}%"
    return str(v)


def extract_xlsx():
    path = RAW / "尊享百万保险责任.xlsx"
    ws = openpyxl.load_workbook(path).active
    merged = {}
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                merged[(r, c)] = (r1, c1)
    cells = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                cells[c.coordinate] = {"value": c.value, "number_format": c.number_format,
                                       "display": cell_text(c)}
    (EXTRACTED / "xlsx_cells.json").write_text(json.dumps(
        {"source": str(path.relative_to(ROOT)), "sheet": ws.title, "dimensions": ws.dimensions,
         "merged_ranges": sorted(str(r) for r in ws.merged_cells.ranges), "cells": cells},
        ensure_ascii=False, indent=1, default=str), encoding="utf-8")

    def val(r, c):
        r0, c0 = merged.get((r, c), (r, c))
        return cell_text(ws.cell(r0, c0))

    items, current = [], None
    for r in range(5, ws.max_row + 1):
        name = val(r, 3)
        # C37 (特定药品费用) is not merged in the source; blank name rows belong to the item above.
        if current is None or (name is not None and name != current["责任名称"]):
            current = {"id": f"R{len(items) + 1:02d}", "责任&加油包": val(r, 1), "人群": val(r, 2),
                       "责任名称": name, "等待期": val(r, 5), "保额": val(r, 6), "赔付比例": [],
                       "免赔额/免赔天数": val(r, 8), "医院范围": [], "责任描述": [],
                       "source_rows": [r, r]}
            items.append(current)
        current["source_rows"][1] = r
        d = ws.cell(r, 4).value
        if d is not None:
            current["责任描述"].append({"cell": f"D{r}", "text": str(d)})
        for key, col in (("赔付比例", 7), ("医院范围", 9)):
            v = val(r, col)
            if v is not None and v not in current[key]:
                current[key].append(v)

    # Every non-empty source cell from row 5 on must be represented verbatim in the structured output.
    blob = json.dumps(items, ensure_ascii=False)
    for coord, info in cells.items():
        row = int(re.sub(r"\D", "", coord))
        if row >= 5:
            assert json.dumps(info["display"], ensure_ascii=False)[1:-1] in blob, f"missing {coord}"
    out = {"product": val(1, 1), "table": val(2, 1),
           "source": str(path.relative_to(ROOT)),
           "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
           "note": "字段值逐字取自源单元格；合并单元格已向下填充；百分比数值按单元格格式显示",
           "items": items}
    (CLEAN / "保险责任.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [f"# {out['product']} {out['table']}", "",
             f"> 来源：`{out['source']}`（sha256 `{out['sha256'][:16]}…`）。内容逐字取自原表，仅做结构化排版。", ""]
    group = None
    for it in items:
        if it["责任&加油包"] != group:
            group = it["责任&加油包"]
            lines += [f"## {group}", ""]
        lines += [f"### [{it['id']}] {it['责任名称']}", "",
                  f"- 人群：{it['人群']}", f"- 等待期：{it['等待期']}", f"- 保额：{it['保额']}",
                  f"- 赔付比例：{'；'.join(it['赔付比例'])}", f"- 免赔额/免赔天数：{it['免赔额/免赔天数']}",
                  f"- 医院范围：{'；'.join(it['医院范围'])}",
                  f"- 原表行：{it['source_rows'][0]}–{it['source_rows'][1]}", "", "责任描述：", ""]
        lines += [f"- `{d['cell']}` {d['text']}" for d in it["责任描述"]]
        lines.append("")
    (CLEAN / "保险责任.md").write_text("\n".join(lines), encoding="utf-8")
    report["xlsx"][path.name] = {"items": len(items), "non_empty_cells": len(cells),
                                 "all_cells_represented": True}


def copy_md():
    src = RAW / "场景及问题说明.md"
    (CLEAN / "场景及问题说明.md").write_bytes(src.read_bytes())
    report["md"][src.name] = {"copied_byte_identical": True}


if __name__ == "__main__":
    extract_pdfs()
    extract_xlsx()
    copy_md()
    (ROOT / "data/extraction_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    sys.path.insert(0, str(ROOT / 'src'))
    from agent.source_locations import INDEX, build_index
    INDEX.parent.mkdir(parents=True, exist_ok=True)
    INDEX.write_text(json.dumps(build_index(), ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    sys.exit(0)
